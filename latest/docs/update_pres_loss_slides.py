"""Replace the E1/E2 loss figures on pres.pptx slides 23-24 with the
matched-objective learning curves, and update their wording.

Both lines now use one common loss -- clean crops, no angular margin, no
label smoothing, cross-entropy over the 30-benchmark gallery -- measured on
the training people (solid) and the held-out validation people (dashed),
from runs/e1e2/learning_curve_matched/. Only those two slides change.

Usage (from docs/):
    python update_pres_loss_slides.py [--pptx pres.pptx]
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Emu  # noqa: E402

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SLIDES = {23: "e1", 24: "e2"}
SUBTITLE = "One common loss: solid = training people, dashed = held-out validation people; one colour per fold."
FOOTNOTE = ("Same objective on both sides (clean crops, no margin, no smoothing, 30-benchmark gallery), "
            "so the gap between the lines is a real generalisation gap. Stars mark the selected epoch, "
            "chosen by validation accuracy.")


def render(exp: str, out_path: Path, wh=(9, 3.6)):
    plt.rcParams.update({"font.family": "Arial", "font.size": 13, "axes.spines.top": False,
                         "axes.spines.right": False, "axes.labelsize": 13, "xtick.labelsize": 12,
                         "ytick.labelsize": 12, "legend.fontsize": 11})
    fig, ax = plt.subplots(figsize=wh, layout="constrained")
    for f in range(1, 7):
        d = json.loads((ROOT / f"runs/e1e2/learning_curve_matched/fold{f}/{exp}_person_model_history.json").read_text())
        h = d["history"]
        for key, style in [("train_loss_matched", "-"), ("val_loss", "--")]:
            valid = [r for r in h if r.get(key) is not None]
            ax.plot([r["epoch"] for r in valid], [r[key] for r in valid], style, color=plt.cm.tab10(f - 1),
                    label=f"Fold {f}" if key == "train_loss_matched" else None, lw=1.7)
        sel = next(r for r in h if r["epoch"] == d["best_epoch"])
        ax.plot(sel["epoch"], sel["val_loss"], "*", color=plt.cm.tab10(f - 1), markersize=11,
                markeredgecolor="black", markeredgewidth=.4, zorder=5)
    ax.set(xlabel="Epoch (0 = pretrained model)", ylabel="Gallery cross-entropy")
    ax.grid(alpha=.2)
    ax.legend(ncol=3, loc="upper right")
    fig.savefig(out_path, dpi=220, facecolor="white")
    plt.close(fig)
    return out_path


def set_text(shape, txt):
    """Overwrite the text, keeping the first run's formatting."""
    p = shape.text_frame.paragraphs[0]
    if not p.runs:
        p.text = txt
        return
    p.runs[0].text = txt
    for r in p.runs[1:]:
        r.text = ""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx", default=str(HERE / "pres.pptx"))
    ap.add_argument("--assets", default=str(HERE / "pres_assets"))
    args = ap.parse_args()

    prs = Presentation(args.pptx)
    assets = Path(args.assets)
    assets.mkdir(parents=True, exist_ok=True)

    for number, exp in SLIDES.items():
        slide = prs.slides[number - 1]
        title = slide.shapes[0].text_frame.text
        assert exp.upper() in title, f"slide {number} is '{title}', not the {exp.upper()} loss slide"
        png = render(exp, assets / f"{exp}_loss_matched.png")

        pics = [sh for sh in slide.shapes if sh.shape_type == 13]
        assert len(pics) == 1, f"slide {number} has {len(pics)} pictures"
        old = pics[0]
        left, top, width, height = old.left, old.top, old.width, old.height
        old._element.getparent().remove(old._element)
        slide.shapes.add_picture(str(png), left, top, width, height)

        boxes = [sh for sh in slide.shapes if sh.has_text_frame]
        set_text(boxes[1], SUBTITLE)                 # subtitle under the title
        set_text(boxes[-1], FOOTNOTE)                # footnote under the figure
        print(f"slide {number}: replaced figure with {png.name} "
              f"({Emu(width).inches:.2f}x{Emu(height).inches:.2f} in) and updated wording")

    prs.save(args.pptx)
    print(f"saved {args.pptx}")


if __name__ == "__main__":
    main()
